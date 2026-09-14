# -*- coding: utf-8 -*-
"""Tests für die Punkteaufschlüsselung (Transparenz des Scorings).

Aufruf aus dem Projekt-Root:  python tests/test_scoring.py

Alle erwarteten Punkte sind von Hand gerechnet (Rechenweg im Kommentar) und
gegen einen festen Stichtag (01.01.2026) geprüft. Nutzt eine eigene
In-Memory-SQLite-Datenbank; die Anwendungsdatenbank bleibt unberührt.
"""
import io
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from backend import models, schemas, score_export, scoring

# Labels enthalten Formelzeichen (≥, ×, Σ, Δ), die die Windows-Konsole sonst nicht kodiert.
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

STICHTAG = pd.Timestamp("2026-01-01")
clock = {"now": STICHTAG}
scoring.system_now = lambda: clock["now"]

failures: list[str] = []


def check(label: str, condition: bool, detail: str = ""):
    if condition:
        print(f"  OK   {label}")
    else:
        print(f"  FAIL {label} {detail}")
        failures.append(label)


def check_close(label: str, actual, expected, tol=1e-6):
    check(label, actual is not None and abs(actual - expected) < tol,
          f"(erwartet {expected!r}, war {actual!r})")


def make_session():
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    models.Base.metadata.create_all(bind=engine)
    db = sessionmaker(bind=engine)()
    scoring.initialize_config(db)
    return db


def add_household(db, name, people, is_resident=False, **kwargs):
    hh = models.Household(name=name, is_resident=is_resident, **kwargs)
    db.add(hh)
    db.flush()
    for person in people:
        db.add(models.Person(**{"household_id": hh.id, "last_name": name, **person}))
    db.flush()
    db.refresh(hh)
    return hh


def build_example(db):
    """Bewohner-Referenz und zwei Bewerber-Haushalte.

    Bewohner (IST-Verteilung):
      R1 35 J., w, Beruf 1, Bildung 7
      R2 41 J., w, Beruf 2, Bildung 7
      R3 66 J., m, Beruf 1, Bildung 6
      R4 11 J., m, Beruf 0, keine Bildung  -> Alter: unter 20; Beruf/Bildung: ohne Angabe
    => Alter (Basis 3): 30-39 1/3, 40-49 1/3, 60-69 1/3
       Geschlecht (Basis 4): w 0,5, m 0,5
       Beruf (Basis 3): 1 2/3, 2 1/3      Bildung (Basis 3): 7 2/3, 6 1/3
    """
    add_household(db, "Bewohner", [
        {"first_name": "R1", "birth_date": datetime(1990, 6, 1), "gender": "w",
         "occupation_type": "1", "education_level": "7"},
        {"first_name": "R2", "birth_date": datetime(1985, 1, 1), "gender": "w",
         "occupation_type": "2", "education_level": "7"},
        {"first_name": "R3", "birth_date": datetime(1960, 1, 1), "gender": "m",
         "occupation_type": "1", "education_level": "6"},
        {"first_name": "R4", "birth_date": datetime(2015, 1, 1), "gender": "m",
         "occupation_type": "0", "education_level": None},
    ], is_resident=True)

    a = add_household(db, "Albers", [
        {"first_name": "Anna", "birth_date": datetime(1998, 3, 1), "gender": "w",
         "occupation_type": "4", "education_level": "4", "member_since": datetime(2012, 1, 1)},
        {"first_name": "Ben", "birth_date": datetime(1996, 1, 1), "gender": "m",
         "occupation_type": "1", "education_level": "7", "member_since": datetime(2020, 1, 1)},
        {"first_name": "Carl", "birth_date": datetime(2020, 1, 1), "gender": "m",
         "occupation_type": "0", "education_level": None},
        # archiviert: das frühe Eintrittsdatum darf nicht zählen
        {"first_name": "Alt", "archived": True, "member_since": datetime(2000, 1, 1)},
    ], engagement_score=0.8, cultural_diversity_score=0.5, special_needs_score=0.0)

    b = add_household(db, "Berg", [
        {"first_name": "Dora", "birth_date": datetime(1950, 1, 1), "gender": "d",
         "occupation_type": "", "education_level": "8", "member_since": datetime(2023, 7, 1)},
    ], engagement_score=0.2)
    db.commit()
    return a, b


def crit(explanation, key):
    return next(c for c in explanation["criteria"] if c["key"] == key)


def term(criterion, group):
    return next((t for t in criterion["terms"] if t["group"] == group), None)


# Erwartete Werte Haushalt Albers -------------------------------------------
# Je Person (Ziel − Ist) / Ziel, summiert über die Personen, × Gewicht.
# Alter:  20-29 (0,24350445 − 0) / 0,24350445 = 1 × 1 Pers. = 1
#         30-39 Ist 1/3 ≥ Ziel 0,17529197 -> 0; Carl (6 J.) nicht berücksichtigt
#         Teilscore 1 × Gewicht 2 = 2
A_AGE = 1.0 * 2
# Geschlecht: w (0,5199 − 0,5) / 0,5199 = 0,038277 × 1; m Ist 0,5 ≥ Ziel 0,4799 -> 0
#         × Gewicht 2 = 0,076553
A_GENDER = 0.0199 / 0.5199 * 2
# Beruf:  4 Ist 0 -> 1 × 1; 1 Ist 2/3 ≥ Ziel -> 0; Carl Kat. 0 -> × Gewicht 1 = 1
A_OCC = 1.0
# Bildung: 4 Ist 0 -> 1; 7 Ist 2/3 ≥ Ziel -> 0 -> × Gewicht 1 = 1
A_EDU = 1.0
A_CULT = 0.5 * 1
A_SPECIAL = 0.0
# Mitgliedsdauer je Person, summiert:
#   Anna seit 01.01.2012 -> 14 × 365 + 4 Schalttage = 5114 Tage = 14,001 J. -> min(14,001; 10) / 10 = 1
#   Ben  seit 01.01.2020 -> 6 × 365 + 2 Schalttage = 2192 Tage = 6,001 J. -> 6,001 / 10 = 0,6001
#   Carl ohne Eintrittsdatum -> nicht berücksichtigt; "Alt" archiviert -> zählt nicht
#   Teilscore 1,6001 × Gewicht 2 = 3,2003
A_MEMBERSHIP = (1.0 + 2192 / 365.25 / 10) * 2
A_ENGAGEMENT = 0.8 * 5
A_BASE = A_AGE + A_GENDER + A_OCC + A_EDU + A_CULT + A_SPECIAL + A_MEMBERSHIP + A_ENGAGEMENT

# Haushalt Berg ---------------------------------------------------------------
# Alter 70-79 Ist 0 -> 1 × 2; Geschlecht d Ist 0 -> 1 × 2; Beruf leer -> 0;
# Bildung 8 Ist 0 -> 1 × 1; Mitglied seit 01.07.2023 -> 915 Tage / 365,25 = 2,505 J. / 10 × 2;
# Engagement 0,2 × 5
B_BASE = 2.0 + 2.0 + 1.0 + 915 / 365.25 / 10 * 2 + 1.0


def test_hand_calculated():
    print("\n== Aufschlüsselung: von Hand gerechnete Werte ==")
    db = make_session()
    a, b = build_example(db)
    config = scoring.get_config_dict(db)
    reference = scoring.calculate_resident_reference(db)
    ea = scoring.explain_household(a, reference, config)
    eb = scoring.explain_household(b, reference, config)

    check_close("Alter", crit(ea, "diversity_age")["points"], A_AGE)
    check_close("Geschlecht", crit(ea, "diversity_gender")["points"], A_GENDER)
    check_close("Beruf", crit(ea, "diversity_occupation")["points"], A_OCC)
    check_close("Bildung", crit(ea, "diversity_education")["points"], A_EDU)
    check_close("Kulturelle Vielfalt", crit(ea, "diversity_cultural")["points"], A_CULT)
    check_close("Besondere Lebenslagen", crit(ea, "diversity_special_needs")["points"], A_SPECIAL)
    check_close("Mitgliedsdauer (je Person, summiert)", crit(ea, "membership")["points"], A_MEMBERSHIP)
    check_close("Engagement", crit(ea, "engagement")["points"], A_ENGAGEMENT)
    check_close("Grundpunktzahl Albers", ea["base_score"], A_BASE)
    check_close("Grundpunktzahl Berg", eb["base_score"], B_BASE)

    age = crit(ea, "diversity_age")
    t = term(age, "20_29")
    check("Term 20-29: Anna, Ist 0/3", t["persons"] == ["Anna Albers"]
          and t["resident_count"] == 0 and t["resident_basis"] == 3, str(t))
    check_close("Term 20-29: maximale Abweichung = 1 Punkt je Person", t["relative_gap"], 1.0)
    check_close("Term Geschlecht w: (Ziel − Ist) / Ziel",
                term(crit(ea, "diversity_gender"), "f")["relative_gap"], 0.0199 / 0.5199)
    t = term(age, "30_39")
    check("Term 30-39: Ist 1/3 ≥ Ziel -> keine Punkte",
          not t["applies"] and t["value"] == 0.0 and t["resident_count"] == 1, str(t))
    check("Kind unter 20 als nicht berücksichtigt ausgewiesen",
          [p["name"] for p in age["ignored_persons"]] == ["Carl Albers"], str(age["ignored_persons"]))
    check("Beruf Kategorie 0 als nicht berücksichtigt ausgewiesen",
          [p["reason"] for p in crit(ea, "diversity_occupation")["ignored_persons"]]
          == [scoring.IGNORED_REASON_LABELS[scoring.MISSING_CATEGORY_0]])

    membership = crit(ea, "membership")
    persons = membership["membership"]["persons"]
    check("Mitgliedsdauer: Anna und Ben, archivierte Person zählt nicht",
          [p["person_name"] for p in persons] == ["Anna Albers", "Ben Albers"], str(persons))
    check("Mitgliedsdauer: Carl ohne Eintrittsdatum ausgewiesen",
          [p["name"] for p in membership["ignored_persons"]] == ["Carl Albers"], str(membership["ignored_persons"]))
    check_close("Anna: Jahre", persons[0]["years"], 5114 / 365.25)
    check_close("Anna: bei maximalen Jahren begrenzt", persons[0]["capped_years"], 10.0)
    check_close("Anna: maximale Jahre erreicht = genau 1 Punkt", persons[0]["value"], 1.0)
    check_close("Ben: anteilig 6,001 von 10 Jahren", persons[1]["value"], 2192 / 365.25 / 10)
    check_close("Mitgliedsdauer: Summe ohne Kappung", membership["subscore"], 1.0 + 2192 / 365.25 / 10)
    check_close("Berg: anteilig 2,505 von 10 Jahren",
                crit(eb, "membership")["subscore"], 915 / 365.25 / 10)


def test_persons_add_up():
    print("\n== Durchmischung: Beiträge mehrerer Personen werden addiert, ohne Kappung ==")
    db = make_session()
    build_example(db)
    paar = add_household(db, "Paar", [
        {"first_name": "E", "gender": "w", "occupation_type": "4"},
        {"first_name": "F", "gender": "m", "occupation_type": "4"},
        {"first_name": "G", "gender": "m", "occupation_type": "9"},
    ])
    db.commit()
    e = scoring.explain_household(paar, scoring.calculate_resident_reference(db), scoring.get_config_dict(db))
    occ = crit(e, "diversity_occupation")
    check_close("Beruf: 2 Personen Handwerk (je 1) + 1 Naturwiss. (1) = 3", occ["subscore"], 3.0)
    check_close("Term Handwerk: 1 × 2 Personen", term(occ, "4")["value"], 2.0)


def test_invariants():
    print("\n== Invarianten ==")
    db = make_session()
    a, b = build_example(db)
    config = scoring.get_config_dict(db)
    reference = scoring.calculate_resident_reference(db)

    for hh in (a, b):
        e = scoring.explain_household(hh, reference, config)
        check_close(f"{hh.name}: Σ Punkte = Grundpunktzahl",
                    sum(c["points"] for c in e["criteria"]), e["base_score"])
        for c in e["criteria"]:
            if c["kind"] == "target":
                check_close(f"{hh.name}/{c['key']}: Σ Terme = Teilscore",
                            sum(t["value"] for t in c["terms"]), c["subscore"])
            check_close(f"{hh.name}/{c['key']}: Punkte = Teilscore × Gewicht",
                        c["points"], c["subscore"] * c["weight"])

    scoring.run_scoring(db)
    check_close("run_scoring speichert dieselbe Grundpunktzahl (Albers)", a.total_score, A_BASE)
    check_close("run_scoring speichert dieselbe Grundpunktzahl (Berg)", b.total_score, B_BASE)

    stats = scoring.calculate_resident_statistics(db)
    mismatches = [
        (c["key"], g["key"]) for c in stats["categories"] if c["basis"] == "person"
        for g in c["groups"]
        if g["count"] != reference["counts"][f"{c['key']}_{g['key']}"] or c["total"] != reference["basis"][c["key"]]
    ]
    check("Ist-Zähler identisch mit der Ist-Statistik", not mismatches, str(mismatches))


def test_stale_and_simulation():
    print("\n== Veraltet-Hinweis und Simulation ==")
    db = make_session()
    a, _ = build_example(db)
    scoring.run_scoring(db)
    config = scoring.get_config_dict(db)
    reference = scoring.calculate_resident_reference(db)
    check("nach Berechnung nicht veraltet",
          not scoring.explain_household(a, reference, config)["is_stale"])

    a.engagement_score = 1.0
    db.commit()
    e = scoring.explain_household(a, reference, config)
    check("Datenänderung ohne Neuberechnung -> veraltet", e["is_stale"])
    check_close("gespeicherter Wert bleibt", e["stored_score"], A_BASE)
    a.engagement_score = 0.8
    db.commit()

    # Simulation: Engagement 0,8 -> 1,0 ergibt Δ = 0,2 × Gewicht 5 = 1,0
    target = schemas.BreakdownTarget(household_id=a.id, size_rooms=3, with_occupancy=True)
    override = {a.id: schemas.ManualOverride(engagement_score=1.0)}
    [sim] = score_export.household_breakdowns(db, [target], override)
    check_close("Simulation: Δ Grundpunktzahl = Δwert × Gewicht", sim["base_score"] - A_BASE, 1.0)
    check("Simulation verändert den Haushalt nicht", a.engagement_score == 0.8)
    check("Simulation meldet nicht fälschlich 'veraltet'", not sim["is_stale"])
    check_close("Simulation: Gesamt ohne Simulation", sim["original_total_score"], A_BASE + 2.0)


def test_reference_date():
    print("\n== Stichtag: Aufschlüsselung reproduziert die gespeicherte Punktzahl ==")
    db = make_session()
    _, b = build_example(db)
    scoring.run_scoring(db)
    check("Stichtag gespeichert", b.score_calculated_at == STICHTAG.to_pydatetime(), str(b.score_calculated_at))

    target = [schemas.BreakdownTarget(household_id=b.id)]
    clock["now"] = pd.Timestamp("2026-07-01")  # ein halbes Jahr später, ohne Neuberechnung
    try:
        [e] = score_export.household_breakdowns(db, target)
        check_close("Aufschlüsselung = gespeicherte Punktzahl trotz verstrichener Zeit", e["base_score"], B_BASE)
        check("verstrichene Zeit allein gilt nicht als veraltet", not e["is_stale"])
        check("Aufschlüsselung nennt den Stichtag der Berechnung",
              e["calculated_at"] == STICHTAG.to_pydatetime(), str(e["calculated_at"]))
        today = scoring.explain_household(b, scoring.calculate_resident_reference(db), scoring.get_config_dict(db))
        # ein halbes Jahr mehr: + 0,5 / 10 × Gewicht 2 = + 0,1
        check("heute gerechnet wäre die Mitgliedsdauer höher", today["base_score"] > B_BASE + 0.09,
              str(today["base_score"]))

        b.people[0].education_level = "7"  # Datenänderung ohne Neuberechnung
        db.commit()
        [e] = score_export.household_breakdowns(db, target)
        check("Datenänderung seit der Berechnung -> veraltet", e["is_stale"])

        scoring.run_scoring(db)
        [e] = score_export.household_breakdowns(db, target)
        check("nach Neuberechnung neuer Stichtag, nicht veraltet",
              not e["is_stale"] and e["calculated_at"] == clock["now"].to_pydatetime())
    finally:
        clock["now"] = STICHTAG


def test_occupancy_and_factors():
    print("\n== Wohnraumausnutzung und Formelparameter ==")
    db = make_session()
    a, _ = build_example(db)

    def breakdown(**kwargs):
        return score_export.household_breakdowns(
            db, [schemas.BreakdownTarget(household_id=a.id, **kwargs)])[0]

    check("ohne Kategorie keine Wohnraumausnutzung", breakdown()["occupancy"] is None)
    check_close("3 Mitglieder, 3 Zimmer -> 1 × 2", breakdown(size_rooms=3, with_occupancy=True)["occupancy"]["points"], 2.0)
    check_close("3 Mitglieder, 4 Zimmer -> 0", breakdown(size_rooms=4, with_occupancy=True)["occupancy"]["points"], 0.0)
    check_close("ohne Zimmerangabe -> erfüllt", breakdown(size_rooms=None, with_occupancy=True)["occupancy"]["points"], 2.0)
    check_close("Gesamt = Grund + Ausnutzung", breakdown(size_rooms=3, with_occupancy=True)["total_score"], A_BASE + 2.0)

    def set_config(key, value):
        db.query(models.ScoringConfig).filter(models.ScoringConfig.key == key).one().value = value
        db.commit()

    check("keine Faktor-Parameter mehr in der Konfiguration",
          not any(k.startswith("factor_") or k == "cap_membership_years" for k in scoring.get_config_dict(db)))
    set_config("max_membership_years", 20.0)
    check_close("Maximum 20 Jahre -> (14,001 + 6,001) / 20 × Gewicht 2", crit(breakdown(), "membership")["points"],
                (5114 + 2192) / 365.25 / 20 * 2)


def test_export():
    print("\n== Excel-Export ==")
    db = make_session()
    a, b = build_example(db)
    request = schemas.BreakdownExportRequest(
        targets=[schemas.BreakdownTarget(household_id=a.id, size_rooms=3, with_occupancy=True),
                 schemas.BreakdownTarget(household_id=b.id, size_rooms=3, with_occupancy=True)],
        overrides={b.id: schemas.ManualOverride(engagement_score=1.0)},
    )
    wb = load_workbook(io.BytesIO(score_export.build_export(db, request)))
    check("Blätter", wb.sheetnames == ["Vergleich", "1 Albers", "2 Berg", "Konfiguration"], str(wb.sheetnames))
    compare = wb["Vergleich"]
    labels = [compare.cell(row=r, column=2).value for r in range(1, compare.max_row + 1)]
    for label in ("Grundpunktzahl", "Gesamtpunktzahl", "Rang in der Auswahl", "Δ durch Simulation"):
        check(f"Vergleich enthält '{label}'", label in labels)
    formulas = [c.value for row in wb["1 Albers"].iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    check("Aufschlüsselung enthält Formeln", len(formulas) > 20, str(len(formulas)))
    sim_values = [c.value for row in wb["2 Berg"].iter_rows() for c in row if c.fill.fgColor.rgb == "00FFF2CC"]
    check("simulierter Wert markiert", 1.0 in sim_values, str(sim_values))


if __name__ == "__main__":
    test_hand_calculated()
    test_persons_add_up()
    test_invariants()
    test_stale_and_simulation()
    test_reference_date()
    test_occupancy_and_factors()
    test_export()
    print()
    if failures:
        print(f"{len(failures)} Fehler: {failures}")
        sys.exit(1)
    print("Alle Tests bestanden.")
