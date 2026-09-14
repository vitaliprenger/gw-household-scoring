"""Punkteaufschluesselung mehrerer Haushalte und ihr Excel-Export.

Die Zahlen stammen aus ``scoring.explain_household`` -- derselben Funktion, die
auch die gespeicherte Grundpunktzahl berechnet. Der Export schreibt die
Eingangswerte als Zahlen und alle Rechenschritte als **Excel-Formeln**, damit
sich jede Punktzahl in Excel nachrechnen laesst.
"""
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from . import models, schemas, scoring

BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDE6F0")
SIMULATED_FILL = PatternFill("solid", fgColor="FFF2CC")
NUMBER = "0.000"
PERCENT = "0.00%"


def household_breakdowns(db: Session, targets: list[schemas.BreakdownTarget],
                         overrides: dict | None = None) -> list[dict]:
    """Aufschluesselung je angefragtem Haushalt, optional mit Wohnraumausnutzung und Simulation.

    Unbekannte Haushalte werden uebersprungen. ``overrides`` bildet
    household_id → ``schemas.ManualOverride`` ab; ``is_stale`` bezieht sich immer
    auf die gespeicherten, nicht auf die simulierten Werte.
    """
    config = scoring.get_config_dict(db)
    reference_cache: dict = {}
    overrides = overrides or {}

    results = []
    for target in targets:
        household = db.get(models.Household, target.household_id)
        if household is None:
            continue
        # Zum Stichtag der letzten Berechnung -- so stimmt die Summe mit der Rangliste ueberein.
        original = scoring.explain_as_calculated(db, household, config, reference_cache)
        override = overrides.get(target.household_id)
        explanation = (
            scoring.explain_as_calculated(db, household, config, reference_cache, override.model_dump())
            if override else dict(original)
        )
        explanation["is_stale"] = original["is_stale"]
        explanation["original_base_score"] = original["base_score"]

        occupancy = None
        if target.with_occupancy:
            occupancy = scoring.explain_occupancy(explanation["member_count"], target.size_rooms, config)
        explanation["occupancy"] = occupancy
        explanation["total_score"] = explanation["base_score"] + (occupancy["points"] if occupancy else 0.0)
        explanation["original_total_score"] = (
            original["base_score"] + (occupancy["points"] if occupancy else 0.0)
        )
        explanation["original_values"] = {
            c["field"]: c["value"] for c in original["criteria"] if c["manual"]
        }
        results.append(explanation)
    return results


# --- Excel ---------------------------------------------------------------------

def _sheet_title(index: int, name: str) -> str:
    clean = re.sub(r"[\[\]:*?/\\']", "", name or "")
    return f"{index} {clean}"[:31]


def _ref(sheet: str, cell: str) -> str:
    return f"='{sheet}'!{cell}"


def _header(ws, row: int, values: list):
    for col, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _write_household_sheet(ws, explanation: dict) -> dict:
    """Schreibt die Aufschluesselung eines Haushalts; liefert die Zellen der Kriterienpunkte."""
    ws["A1"] = explanation["name"]
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Mitglieder"
    ws["B2"] = explanation["member_count"]
    ws["A3"] = "Stichtag (letzte Berechnung)"
    ws["B3"] = explanation["calculated_at"]
    ws["B3"].number_format = "DD.MM.YYYY HH:MM"
    ws["A4"] = "Gespeicherte Grundpunktzahl (letzte Berechnung)"
    ws["B4"] = explanation["stored_score"]
    ws["B4"].number_format = NUMBER

    row = 6
    points_cells: dict[str, str] = {}
    for criterion in explanation["criteria"]:
        ws.cell(row=row, column=1, value=f"{criterion['category']} – {criterion['label']}").font = Font(bold=True, size=12)
        row += 1

        if criterion["kind"] == "target":
            _header(ws, row, ["Gruppe", "Personen", "Anzahl", "Ziel", "Bewohner der Gruppe",
                              "Bewohner mit Angabe", "Ist = E / F", "Ziel − Ist",
                              "je Person = WENN(Ist < Ziel; (Ziel − Ist) / Ziel; 0)",
                              "Wert = je Person × Anzahl"])
            row += 1
            first = row
            for term in criterion["terms"]:
                ws.cell(row=row, column=1, value=term["label"])
                ws.cell(row=row, column=2, value=", ".join(term["persons"]))
                ws.cell(row=row, column=3, value=term["count"])
                ws.cell(row=row, column=4, value=term["target"]).number_format = PERCENT
                ws.cell(row=row, column=5, value=term["resident_count"])
                ws.cell(row=row, column=6, value=term["resident_basis"])
                ws.cell(row=row, column=7, value=f"=IF(F{row}>0,E{row}/F{row},0)").number_format = PERCENT
                ws.cell(row=row, column=8, value=f"=D{row}-G{row}").number_format = PERCENT
                ws.cell(row=row, column=9,
                        value=f"=IF(G{row}<D{row},H{row}/D{row},0)").number_format = NUMBER
                ws.cell(row=row, column=10, value=f"=I{row}*C{row}").number_format = NUMBER
                row += 1
            last = row - 1
            if criterion["ignored_persons"]:
                ws.cell(row=row, column=1, value="Nicht berücksichtigt")
                ws.cell(row=row, column=2, value="; ".join(
                    f"{p['name']} ({p['reason']})" for p in criterion["ignored_persons"]))
                row += 1
            ws.cell(row=row, column=9, value="Teilscore")
            subscore = f"J{row}"
            ws[subscore] = f"=SUM(J{first}:J{last})" if criterion["terms"] else 0
            ws[subscore].number_format = NUMBER
            row += 1

        elif criterion["kind"] == "membership":
            m = criterion["membership"]
            ws.cell(row=row, column=1, value="Maximale Mitgliedsjahre")
            ws.cell(row=row, column=2, value=m["max_years"])
            maximum = f"$B${row}"
            row += 1
            _header(ws, row, ["Person", "Mitglied seit", "Jahre = GANZZAHL(Stichtag − Eintritt) / 365,25",
                              "Wert = MIN(Jahre; Maximum) / Maximum"])
            row += 1
            first = row
            for person in m["persons"]:
                ws.cell(row=row, column=1, value=person["person_name"])
                ws.cell(row=row, column=2, value=person["member_since"]).number_format = "DD.MM.YYYY"
                ws.cell(row=row, column=3, value=f"=INT($B$3-B{row})/365.25").number_format = NUMBER
                ws.cell(row=row, column=4,
                        value=f"=MAX(0,MIN(C{row},{maximum}))/{maximum}").number_format = NUMBER
                row += 1
            last = row - 1
            if criterion["ignored_persons"]:
                ws.cell(row=row, column=1, value="Nicht berücksichtigt")
                ws.cell(row=row, column=2, value="; ".join(
                    f"{p['name']} ({p['reason']})" for p in criterion["ignored_persons"]))
                row += 1
            ws.cell(row=row, column=1, value="Teilscore = Summe der Werte")
            ws.cell(row=row, column=2,
                    value=f"=SUM(D{first}:D{last})" if m["persons"] else 0).number_format = NUMBER
            subscore = f"B{row}"
            row += 1

        else:  # manuell
            simulated = explanation["original_values"].get(criterion["field"]) != criterion["value"]
            ws.cell(row=row, column=1, value="Erfüllungsgrad (0–1, manuell bewertet)"
                    + (" – simuliert" if simulated else ""))
            cell = ws.cell(row=row, column=2, value=criterion["value"])
            if simulated:
                cell.fill = SIMULATED_FILL
                ws.cell(row=row, column=3,
                        value=f"gespeichert: {explanation['original_values'].get(criterion['field'])}")
            subscore = f"B{row}"
            row += 1

        label_col = 9 if criterion["kind"] == "target" else 1
        value_col = "J" if criterion["kind"] == "target" else "B"
        ws.cell(row=row, column=label_col, value="Gewicht")
        ws[f"{value_col}{row}"] = criterion["weight"]
        weight = f"{value_col}{row}"
        row += 1
        ws.cell(row=row, column=label_col, value="Punkte = Teilscore × Gewicht").font = BOLD
        ws[f"{value_col}{row}"] = f"={subscore}*{weight}"
        ws[f"{value_col}{row}"].number_format = NUMBER
        ws[f"{value_col}{row}"].font = BOLD
        points_cells[criterion["key"]] = f"{value_col}{row}"
        row += 2

    ws.cell(row=row, column=1, value="Grundpunktzahl = Summe der Punkte").font = BOLD
    ws.cell(row=row, column=2, value="=" + "+".join(points_cells.values())).number_format = NUMBER
    ws.cell(row=row, column=2).font = BOLD
    base_cell = f"B{row}"
    row += 2

    occupancy_cell = None
    occ = explanation["occupancy"]
    if occ:
        ws.cell(row=row, column=1, value="Wohnraumausnutzung").font = Font(bold=True, size=12)
        row += 1
        ws.cell(row=row, column=1, value="Mitglieder")
        ws.cell(row=row, column=2, value=occ["members"])
        members = f"B{row}"
        row += 1
        ws.cell(row=row, column=1, value="Zimmer (leer = ohne Zimmerangabe)")
        ws.cell(row=row, column=2, value=occ["size_rooms"])
        rooms = f"B{row}"
        row += 1
        ws.cell(row=row, column=1, value="Erfüllt = WENN(ohne Zimmer; 1; Mitglieder ≥ Zimmer)")
        ws.cell(row=row, column=2, value=f'=IF({rooms}="",1,IF({members}>={rooms},1,0))')
        fulfilled = f"B{row}"
        row += 1
        ws.cell(row=row, column=1, value="Gewicht")
        ws.cell(row=row, column=2, value=occ["weight"])
        weight = f"B{row}"
        row += 1
        ws.cell(row=row, column=1, value="Punkte = Erfüllt × Gewicht").font = BOLD
        ws.cell(row=row, column=2, value=f"={fulfilled}*{weight}").number_format = NUMBER
        occupancy_cell = f"B{row}"
        row += 2

    ws.cell(row=row, column=1, value="Gesamtpunktzahl").font = Font(bold=True, size=12)
    total = f"={base_cell}+{occupancy_cell}" if occupancy_cell else f"={base_cell}"
    ws.cell(row=row, column=2, value=total).number_format = NUMBER
    ws.cell(row=row, column=2).font = Font(bold=True, size=12)

    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 30
    for col in range(3, 11):
        ws.column_dimensions[get_column_letter(col)].width = 16

    return {"criteria": points_cells, "base": base_cell, "occupancy": occupancy_cell, "total": f"B{row}"}


def _write_config_sheet(ws, db: Session):
    ws["A1"] = "Bewertungskonfiguration zum Zeitpunkt des Exports"
    ws["A1"].font = BOLD
    _header(ws, 3, ["Schlüssel", "Beschreibung", "Wert"])
    for row, conf in enumerate(
        sorted(db.query(models.ScoringConfig).all(), key=lambda c: c.key), 4
    ):
        ws.cell(row=row, column=1, value=conf.key)
        ws.cell(row=row, column=2, value=conf.description)
        ws.cell(row=row, column=3, value=conf.value)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 14


def build_export(db: Session, request: schemas.BreakdownExportRequest) -> bytes:
    """Excel-Arbeitsmappe: Vergleich, je Haushalt eine Aufschluesselung, Konfiguration."""
    explanations = household_breakdowns(db, request.targets, request.overrides)

    wb = Workbook()
    compare = wb.active
    compare.title = "Vergleich"
    _write_config_sheet(wb.create_sheet("Konfiguration"), db)

    cells = []
    for index, explanation in enumerate(explanations, 1):
        title = _sheet_title(index, explanation["name"])
        ws = wb.create_sheet(title, index)
        cells.append((title, _write_household_sheet(ws, explanation)))

    compare["A1"] = "Punktevergleich"
    compare["A1"].font = Font(bold=True, size=14)
    compare["A2"] = "Exportiert am"
    compare["B2"] = scoring.system_now().to_pydatetime()
    compare["B2"].number_format = "DD.MM.YYYY HH:MM"
    if request.overrides:
        compare["A3"] = "Gelb markiert: simulierte manuelle Bewertung (nicht gespeichert)"
        compare["A3"].fill = SIMULATED_FILL

    row = 5
    _header(compare, row, ["Kategorie", "Kriterium"] + [e["name"] for e in explanations])
    row += 1
    first_col = 3
    last_col = first_col + len(explanations) - 1

    if explanations:
        for criterion in explanations[0]["criteria"]:
            compare.cell(row=row, column=1, value=criterion["category"])
            compare.cell(row=row, column=2, value=criterion["label"] + (" (manuell)" if criterion["manual"] else ""))
            for offset, (title, sheet_cells) in enumerate(cells):
                cell = compare.cell(row=row, column=first_col + offset,
                                    value=_ref(title, sheet_cells["criteria"][criterion["key"]]))
                cell.number_format = NUMBER
                expl = explanations[offset]
                crit = next(c for c in expl["criteria"] if c["key"] == criterion["key"])
                if crit["manual"] and expl["original_values"].get(crit["field"]) != crit["value"]:
                    cell.fill = SIMULATED_FILL
            row += 1

    def summary_row(label: str, key: str, bold=False):
        nonlocal row
        compare.cell(row=row, column=2, value=label).font = Font(bold=bold)
        for offset, (title, sheet_cells) in enumerate(cells):
            value = _ref(title, sheet_cells[key]) if sheet_cells[key] else 0
            cell = compare.cell(row=row, column=first_col + offset, value=value)
            cell.number_format = NUMBER
            cell.font = Font(bold=bold)
        row += 1
        return row - 1

    summary_row("Grundpunktzahl", "base", bold=True)
    summary_row("Wohnraumausnutzung", "occupancy")
    total_row = summary_row("Gesamtpunktzahl", "total", bold=True)

    if explanations:
        span = f"${get_column_letter(first_col)}${total_row}:${get_column_letter(last_col)}${total_row}"
        compare.cell(row=row, column=2, value="Rang in der Auswahl")
        for col in range(first_col, last_col + 1):
            compare.cell(row=row, column=col, value=f"=RANK({get_column_letter(col)}{total_row},{span})")
        row += 1
        compare.cell(row=row, column=2, value="Abstand zum Ersten")
        for col in range(first_col, last_col + 1):
            compare.cell(row=row, column=col,
                         value=f"=MAX({span})-{get_column_letter(col)}{total_row}").number_format = NUMBER
        row += 1
        if request.overrides:
            compare.cell(row=row, column=2, value="Gesamtpunktzahl ohne Simulation")
            for offset, explanation in enumerate(explanations):
                compare.cell(row=row, column=first_col + offset,
                             value=explanation["original_total_score"]).number_format = NUMBER
            row += 1
            compare.cell(row=row, column=2, value="Δ durch Simulation")
            for col in range(first_col, last_col + 1):
                letter = get_column_letter(col)
                compare.cell(row=row, column=col,
                             value=f"={letter}{total_row}-{letter}{row - 1}").number_format = NUMBER
            row += 1

    compare.column_dimensions["A"].width = 18
    compare.column_dimensions["B"].width = 34
    for col in range(first_col, last_col + 1):
        compare.column_dimensions[get_column_letter(col)].width = 22

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
